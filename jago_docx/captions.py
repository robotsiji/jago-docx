"""
Captions engine for Jago Docx.
Strictly places captions DIRECTLY UNDERNEATH tables and figures.
Utilizes native Word SEQ fields (SEQ Table, SEQ Figure) to enable dynamic numbering
and integration with the auto-generated List of Figures and List of Tables.
"""

import os
from typing import List, Optional, Union
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from .typography import COLOR_TEXT_PRIMARY, COLOR_MUTED, COLOR_HEADING_1

W_NS = nsdecls("w")


def create_caption_paragraph(doc, label: str, seq_identifier: str, caption_text: str, align=WD_ALIGN_PARAGRAPH.CENTER):
    """
    Creates a caption paragraph located right under the respective visual element.
    Uses Word SEQ field: `<w:fldSimple w:instr="SEQ {seq_identifier} \\* ARABIC"/>`.
    """
    safe_caption = caption_text.strip()

    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15

    try:
        p.style = doc.styles["Caption"]
    except Exception:
        pass

    # Run 1: Label (e.g. Table / Figure)
    r_label = p.add_run(f"{label} ")
    r_label.font.name = "Arial"
    r_label.font.bold = True
    r_label.font.size = Pt(9.5)
    r_label.font.color.rgb = COLOR_MUTED

    # SEQ Field
    fld = parse_xml(
        f'<w:fldSimple {W_NS} w:instr="SEQ {seq_identifier} \\* ARABIC">'
        f'<w:r><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:b/><w:sz w:val="19"/><w:color w:val="475569"/></w:rPr>'
        f'<w:t>1</w:t></w:r></w:fldSimple>'
    )
    p._p.append(fld)

    # Run 2: Description
    r_desc = p.add_run(f": {safe_caption}")
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = COLOR_MUTED

    return p


def _set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """
    Set clean cell padding in dxa (1 pt = 20 dxa).
    Default 6pt top/bottom, 8pt left/right.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f"""
    <w:tcMar {W_NS}>
      <w:top w:w="{top}" w:type="dxa"/>
      <w:bottom w:w="{bottom}" w:type="dxa"/>
      <w:left w:w="{left}" w:type="dxa"/>
      <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>
    """)
    tcPr.append(tcMar)


def _set_cell_background(cell, hex_color="F8FAFC"):
    """
    Set background shading on a table cell.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {W_NS} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_table_borders(table, border_color="E2E8F0"):
    """
    Applies clean modern border styling: light horizontal lines, no vertical lines.
    """
    tblPr = table._tbl.tblPr
    borders = parse_xml(f"""
    <w:tblBorders {W_NS}>
      <w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
      <w:left w:val="none"/>
      <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
      <w:right w:val="none"/>
      <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
      <w:insideV w:val="none"/>
    </w:tblBorders>
    """)
    tblPr.append(borders)


def add_table_with_caption(
    doc,
    headers: List[str],
    rows: List[List[Union[str, int, float]]],
    caption: str,
    col_widths: Optional[List[float]] = None,
    align=WD_TABLE_ALIGNMENT.CENTER
):
    """
    Renders a beautifully styled antislop data table,
    then strictly places its caption directly underneath it.
    """
    num_cols = len(headers)
    num_rows = len(rows) + 1

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = align
    _set_table_borders(table)

    # Style Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        _set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        _set_cell_background(cell, "F8FAFC")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(str(title))
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_HEADING_1

    # Style Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            _set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEXT_PRIMARY

    # Apply column widths if specified
    if col_widths and len(col_widths) == num_cols:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # STRICT RULE: CAPTION RIGHT UNDER TABLE
    create_caption_paragraph(doc, label="Table", seq_identifier="Table", caption_text=caption)
    return table


def add_figure_with_caption(
    doc,
    image_path_or_text: str,
    caption: str,
    width_inches: float = 5.0
):
    """
    Renders an image or visual placeholder,
    then strictly places its caption directly underneath it.
    """
    p_fig = doc.add_paragraph()
    p_fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fig.paragraph_format.space_before = Pt(8)
    p_fig.paragraph_format.space_after = Pt(4)
    p_fig.paragraph_format.keep_with_next = True

    if os.path.exists(image_path_or_text):
        p_fig.add_run().add_picture(image_path_or_text, width=Inches(width_inches))
    else:
        # Visual placeholder container
        run_box = p_fig.add_run(f"[{image_path_or_text.strip()}]")
        run_box.font.name = "Arial"
        run_box.font.size = Pt(10)
        run_box.font.italic = True
        run_box.font.color.rgb = COLOR_MUTED

    # STRICT RULE: CAPTION RIGHT UNDER FIGURE
    create_caption_paragraph(doc, label="Figure", seq_identifier="Figure", caption_text=caption)
    return p_fig
