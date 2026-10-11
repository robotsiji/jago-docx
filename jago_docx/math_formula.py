"""
Native editable formula and equation placeholder engine for Jago Docx.
Leverages Office Math Markup Language (OMML / m:oMath) so equations remain
directly editable inside Microsoft Word's Equation Editor.
"""

from typing import Optional
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .typography import COLOR_TEXT_PRIMARY, COLOR_MUTED

MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
W_NS = nsdecls("w")


# Preset formula templates that are fully editable in Word
FORMULA_PRESETS = {
    "linear": "y = m x + b",
    "quadratic": "x = (-b ± √(b² - 4ac)) / (2a)",
    "pythagorean": "c = √(a² + b²)",
    "energy": "E = m c²",
    "fraction": "f(x) = (a + b) / (c + d)",
    "bayes": "P(A|B) = (P(B|A) P(A)) / P(B)",
    "placeholder": "[ Click to edit formula: f(x) = ax² + bx + c ]",
}


def add_editable_formula(
    doc,
    formula_text: Optional[str] = None,
    equation_number: Optional[int] = None,
    preset: Optional[str] = None,
    description: Optional[str] = None
):
    """
    Inserts a native, editable Word equation block (OMML).
    Allows user to edit variables, symbols, and operators in Word Equation Editor.
    
    If equation_number is provided, aligns formula centered with (n) right-aligned.
    If description is provided, adds a small explanatory caption/note underneath.
    """
    # Determine the equation string
    if formula_text and formula_text.strip():
        raw_eq = formula_text.strip()
    elif preset and preset in FORMULA_PRESETS:
        raw_eq = FORMULA_PRESETS[preset]
    else:
        raw_eq = FORMULA_PRESETS["placeholder"]

    # Escape XML special characters
    safe_eq = (
        raw_eq
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )

    if equation_number is not None:
        # Paragraph with centered tab stop (3.25 in) and right tab stop (6.5 in)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8 if not description else 2)
        p.paragraph_format.keep_with_next = True if description else False
        
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
        tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

        # First tab moves to center
        r_tab1 = p.add_run("\t")
        
        # Insert OMML inline element
        omml_xml = f"""
        <m:oMath {MATH_NS}>
          <m:r>
            <m:rPr><m:sty m:val="p"/></m:rPr>
            <m:t xml:space="preserve">{safe_eq}</m:t>
          </m:r>
        </m:oMath>
        """
        p._p.append(parse_xml(omml_xml))

        # Second tab moves to right margin for equation number
        r_tab2 = p.add_run("\t")
        r_num = p.add_run(f"({equation_number})")
        r_num.font.name = "Arial"
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = COLOR_MUTED
    else:
        # Centered standalone display equation
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8 if not description else 2)
        p.paragraph_format.keep_with_next = True if description else False

        omml_para_xml = f"""
        <m:oMathPara {MATH_NS}>
          <m:oMath>
            <m:r>
              <m:rPr><m:sty m:val="p"/></m:rPr>
              <m:t xml:space="preserve">{safe_eq}</m:t>
            </m:r>
          </m:oMath>
        </m:oMathPara>
        """
        p._p.append(parse_xml(omml_para_xml))

    # Optional description / legend underneath
    if description:
        p_desc = doc.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(8)
        r_desc = p_desc.add_run(description.strip())
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(9.5)
        r_desc.font.italic = True
        r_desc.font.color.rgb = COLOR_MUTED

    return p
