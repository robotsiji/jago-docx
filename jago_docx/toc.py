"""
Table of Contents, List of Figures, and List of Tables generator for Jago Docx.
Inserts native OpenXML field codes and enables automatic field updating when opened in Word.
"""

from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .typography import COLOR_HEADING_1, COLOR_MUTED


def enable_update_fields_on_open(doc):
    """
    Instructs Microsoft Word to prompt/update all fields (TOC, page numbers, figure lists)
    automatically upon opening the document.
    """
    settings = doc.settings.element
    # Check if updateFields already exists
    existing = settings.findall(f'{{{nsdecls("w").split("=")[1].strip(chr(34))}}}updateFields')
    if not existing:
        update_elem = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
        settings.append(update_elem)


def _create_complex_toc_field(instr_text: str, fallback_text: str):
    """
    Creates a robust complex field element for TOCs with fallback run.
    """
    w_ns = nsdecls("w")
    xml_str = f"""
    <w:p {w_ns}>
      <w:pPr>
        <w:spacing w:before="0" w:after="120" w:line="276" w:lineRule="auto"/>
      </w:pPr>
      <w:r>
        <w:fldChar w:fldCharType="begin"/>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>
          <w:color w:val="64748B"/>
          <w:sz w:val="20"/>
        </w:rPr>
        <w:instrText xml:space="preserve"> {instr_text} </w:instrText>
      </w:r>
      <w:r>
        <w:fldChar w:fldCharType="separate"/>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>
          <w:color w:val="64748B"/>
          <w:i/>
          <w:sz w:val="19"/>
        </w:rPr>
        <w:t xml:space="preserve">{fallback_text}</w:t>
      </w:r>
      <w:r>
        <w:fldChar w:fldCharType="end"/>
      </w:r>
    </w:p>
    """
    return parse_xml(xml_str)


def add_table_of_contents_page(doc, title: str = "Table of Contents"):
    """
    Auto-generates a dedicated page for the document Table of Contents (TOC).
    Inserts native Word TOC field and appends a page break.
    """
    enable_update_fields_on_open(doc)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.keep_with_next = True

    run = p_title.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING_1

    # Insert complex TOC field for Heading levels 1-3
    toc_p = _create_complex_toc_field(
        'TOC \\o "1-3" \\h \\z \\u',
        "[Table of Contents: Automatically refreshed by Word on open. Press F9 to update anytime]"
    )
    doc._body._element.append(toc_p)
    doc.add_page_break()


def add_list_of_figures_page(doc, title: str = "List of Figures"):
    """
    Auto-generates a dedicated page for the List of Figures.
    Collects captions tagged with SEQ Figure.
    """
    enable_update_fields_on_open(doc)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.keep_with_next = True

    run = p_title.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING_1

    lof_p = _create_complex_toc_field(
        'TOC \\c "Figure" \\h \\z',
        "[List of Figures: Automatically populated by Word. Press F9 to update]"
    )
    doc._body._element.append(lof_p)
    doc.add_page_break()


def add_list_of_tables_page(doc, title: str = "List of Tables"):
    """
    Auto-generates a dedicated page for the List of Tables.
    Collects captions tagged with SEQ Table.
    """
    enable_update_fields_on_open(doc)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.keep_with_next = True

    run = p_title.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING_1

    lot_p = _create_complex_toc_field(
        'TOC \\c "Table" \\h \\z',
        "[List of Tables: Automatically populated by Word. Press F9 to update]"
    )
    doc._body._element.append(lot_p)
    doc.add_page_break()
