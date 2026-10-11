"""
Page number engine for Jago Docx.
Injects native, editable page number placeholders into the document footer
using standard Word field codes (PAGE and NUMPAGES).
"""

from typing import Optional
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from .typography import COLOR_MUTED

W_NS = nsdecls("w")


def setup_editable_page_number(
    doc,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
    format_type: str = "page_of_total",  # 'page_of_total', 'simple_page', 'custom'
    prefix: str = "Page ",
    separator: str = " of ",
    different_first_page: bool = False
):
    """
    Configures footer page numbers across all sections.
    Inserts standard Word field codes (PAGE, NUMPAGES) within a native footer paragraph
    so the user can edit, re-align, or re-style the placeholder directly in Word.
    """
    for section in doc.sections:
        if different_first_page:
            section.different_first_page_header_footer = True
            
        footer = section.footer
        p = footer.paragraphs[0]
        p.text = ""  # Clear any default text
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        # Style font properties
        r_pr_xml = """
        <w:rPr>
          <w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>
          <w:sz w:val="18"/>
          <w:color w:val="64748B"/>
        </w:rPr>
        """

        if format_type == "page_of_total":
            # "Page " run
            r1 = p.add_run(prefix)
            r1.font.name = "Arial"
            r1.font.size = Pt(9)
            r1.font.color.rgb = COLOR_MUTED

            # [PAGE] field
            fld_page = parse_xml(f"""
            <w:fldSimple {W_NS} w:instr="PAGE">
              <w:r>{r_pr_xml}<w:t>1</w:t></w:r>
            </w:fldSimple>
            """)
            p._p.append(fld_page)

            # " of " run
            r2 = p.add_run(separator)
            r2.font.name = "Arial"
            r2.font.size = Pt(9)
            r2.font.color.rgb = COLOR_MUTED

            # [NUMPAGES] field
            fld_total = parse_xml(f"""
            <w:fldSimple {W_NS} w:instr="NUMPAGES">
              <w:r>{r_pr_xml}<w:t>1</w:t></w:r>
            </w:fldSimple>
            """)
            p._p.append(fld_total)

        elif format_type == "simple_page":
            # Just the number, e.g. "1"
            fld_page = parse_xml(f"""
            <w:fldSimple {W_NS} w:instr="PAGE">
              <w:r>{r_pr_xml}<w:t>1</w:t></w:r>
            </w:fldSimple>
            """)
            p._p.append(fld_page)

        elif format_type == "custom":
            # Basic editable placeholder string with PAGE field
            r1 = p.add_run(prefix)
            r1.font.name = "Arial"
            r1.font.size = Pt(9)
            r1.font.color.rgb = COLOR_MUTED

            fld_page = parse_xml(f"""
            <w:fldSimple {W_NS} w:instr="PAGE">
              <w:r>{r_pr_xml}<w:t>1</w:t></w:r>
            </w:fldSimple>
            """)
            p._p.append(fld_page)
